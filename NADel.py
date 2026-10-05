import streamlit as st
from openpyxl import load_workbook
from io import BytesIO

st.set_page_config(page_title="Remove N/A Rows", layout="centered")

st.title("Remove Rows with N/A in Column AZ")

uploaded_file = st.file_uploader(
    "Upload Excel File",
    type=["xlsx", "xlsm"]
)

if uploaded_file:
    try:
        # Load workbook (preserves VBA for xlsm files)
        wb = load_workbook(uploaded_file, keep_vba=True)
        ws = wb.active

        rows_removed = 0

        # Loop from bottom to top so row deletions don't affect indexing
        for row in range(ws.max_row, 1, -1):  # Skip header row
            value = ws.cell(row=row, column=52).value  # AZ = column 52

            if value is None or str(value).strip().upper() == "#N/A":
                ws.delete_rows(row)
                rows_removed += 1

        st.success(f"Removed {rows_removed} row(s).")

        # Save workbook to memory
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Keep original extension
        original_name = uploaded_file.name

        st.download_button(
            label="📥 Download Cleaned File",
            data=output,
            file_name=f"cleaned_{original_name}",
            mime="application/vnd.ms-excel.sheet.macroEnabled.12"
            if original_name.lower().endswith(".xlsm")
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        st.error(f"Error processing file: {e}")